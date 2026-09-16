from pathlib import Path
from openpyxl import load_workbook
import sys

root = Path(__file__).resolve().parent
needles = [s.lower() for s in sys.argv[1:]]
for path in sorted(root.rglob("*.xlsx")):
    rel = str(path.relative_to(root))
    if needles and not any(rel.lower().startswith(n) for n in needles):
        continue
    print(f"\n### {path.relative_to(root)}")
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"-- {ws.title} [{ws.max_row}x{ws.max_column}]")
        if not ws.max_row:
            continue
        for r in range(1, min(ws.max_row, 220) + 1):
            a = ws.cell(r, 1).value
            b = ws.cell(r, 2).value
            if a is not None or b is not None:
                a_s = str(a).replace("\n", " ")[:95] if a is not None else ""
                b_s = str(b).replace("\n", " ")[:55] if b is not None else ""
                print(f"{r:>3}: {a_s} | {b_s}")
