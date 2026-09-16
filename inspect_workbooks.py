from pathlib import Path
from openpyxl import load_workbook
import json

root = Path(__file__).resolve().parent
out = []
for path in sorted(root.rglob("*.xlsx")):
    try:
        wb_formula = load_workbook(path, read_only=True, data_only=False)
        wb_values = load_workbook(path, read_only=True, data_only=True)
        book = {"path": str(path.relative_to(root)), "sheets": []}
        for ws_f in wb_formula.worksheets:
            ws_v = wb_values[ws_f.title]
            rows = []
            for r_idx, (rf, rv) in enumerate(zip(ws_f.iter_rows(), ws_v.iter_rows()), start=1):
                vals = []
                for cf, cv in zip(rf, rv):
                    value = cv.value if cv.value is not None else cf.value
                    if value is not None:
                        vals.append([cf.column, value])
                if vals:
                    rows.append([r_idx, vals])
                if len(rows) >= 35:
                    break
            book["sheets"].append({
                "title": ws_f.title,
                "max_row": ws_f.max_row,
                "max_column": ws_f.max_column,
                "sample_rows": rows,
            })
        out.append(book)
    except Exception as exc:
        out.append({"path": str(path.relative_to(root)), "error": repr(exc)})

(root / "workbook_inventory.json").write_text(json.dumps(out, indent=2, default=str), encoding="utf-8")
print(f"Wrote workbook_inventory.json for {len(out)} workbooks")
