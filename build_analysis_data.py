from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime
import csv
import json
import math
import statistics

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)


def workbook(prefix):
    path = next(ROOT.glob(prefix + "*/*.xlsx"))
    wb = load_workbook(path, read_only=True, data_only=True)
    return path, wb[wb.sheetnames[0]]


def annual_row(prefix, row, header_row=4, start_col=3):
    path, ws = workbook(prefix)
    values = {}
    for col in range(start_col, (ws.max_column or 0) + 1):
        header = ws.cell(header_row, col).value
        if isinstance(header, datetime):
            year = header.year
        elif isinstance(header, str) and header.startswith("FY ") and "Est" not in header:
            try:
                year = int(header.split()[1])
            except ValueError:
                continue
        else:
            continue
        if 2007 <= year <= 2026:
            val = ws.cell(row, col).value
            if isinstance(val, (int, float)) and math.isfinite(val):
                values[year] = float(val)
    return values


def annual_row_multi(prefix, row):
    return annual_row(prefix, row, header_row=4, start_col=5)


def price_series(prefix):
    path, ws = workbook(prefix)
    points = []
    for r in range(8, (ws.max_row or 0) + 1):
        d, p = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(d, datetime) and isinstance(p, (int, float)):
            points.append({"date": d.strftime("%Y-%m-%d"), "price": float(p)})
    return sorted(points, key=lambda x: x["date"])


years = list(range(2007, 2027))
revenue = annual_row("011", 6)
op_income = annual_row("011", 21)
net_income = annual_row("011", 49)
op_margin = annual_row("016", 15)
net_margin = annual_row("016", 19)
roe = annual_row("016", 7)
revenue_growth = annual_row("017", 7)
headcount = annual_row("026", 6)
sales_employee = annual_row("026", 8)
cfo = annual_row("019", 18)
dividends_paid = {y: abs(v) for y, v in annual_row("019", 40).items()}
repurchase_raw = annual_row("019", 45)
buybacks = {y: max(0.0, -v) for y, v in repurchase_raw.items()}
fcf = annual_row("019", 63)

geography = {
    "North America": annual_row_multi("028", 7),
    "Europe": annual_row_multi("028", 9),
    "India": annual_row_multi("028", 11),
    "Rest of World": annual_row_multi("028", 13),
}

segments = {
    "Financial Services": annual_row_multi("030", 7),
    "Retail": annual_row_multi("030", 17),
    "Communication": annual_row_multi("030", 27),
    "Energy & Utilities": annual_row_multi("030", 37),
    "Manufacturing": annual_row_multi("030", 47),
    "Hi-Tech": annual_row_multi("030", 57),
    "Life Sciences": annual_row_multi("030", 67),
    "Others": annual_row_multi("030", 77),
}

tcs_growth = annual_row("048", 7)
tcs_margin = annual_row("047", 15)
tcs_roe = annual_row("047", 7)

infosys_price = price_series("033")
tcs_price = price_series("045")


def cagr(first, last, periods):
    return (last / first) ** (1 / periods) - 1


def pct(a, b):
    return (b / a - 1) * 100


metrics = {
    "revenue_cagr_pct": cagr(revenue[2007], revenue[2026], 19) * 100,
    "net_income_cagr_pct": cagr(net_income[2007], net_income[2026], 19) * 100,
    "fcf_cagr_pct": cagr(fcf[2007], fcf[2026], 19) * 100,
    "headcount_cagr_pct": cagr(headcount[2007], headcount[2026], 19) * 100,
    "sales_per_employee_cagr_pct": cagr(sales_employee[2007], sales_employee[2026], 19) * 100,
    "operating_margin_change_pp": op_margin[2026] - op_margin[2007],
    "net_margin_change_pp": net_margin[2026] - net_margin[2007],
    "roe_change_pp": roe[2026] - roe[2007],
    "headcount_peak_year": max(headcount, key=headcount.get),
    "headcount_peak": max(headcount.values()),
    "headcount_change_2023_2024_pct": pct(headcount[2023], headcount[2024]),
    "sales_per_employee_change_2023_2026_pct": pct(sales_employee[2023], sales_employee[2026]),
    "fcf_conversion_2026_pct": fcf[2026] / net_income[2026] * 100,
    "capital_returns_2026_inr_mn": dividends_paid[2026] + buybacks[2026],
    "north_america_share_change_pp_2009_2026": geography["North America"][2026] - geography["North America"][2009],
    "europe_share_change_pp_2009_2026": geography["Europe"][2026] - geography["Europe"][2009],
    "financial_services_change_pp_2018_2026": segments["Financial Services"][2026] - segments["Financial Services"][2018],
    "manufacturing_change_pp_2018_2026": segments["Manufacturing"][2026] - segments["Manufacturing"][2018],
}

data = {
    "units": {"financials": "INR million", "margins": "percent", "sales_per_employee": "INR"},
    "annual": {
        "revenue": revenue,
        "operating_income": op_income,
        "net_income": net_income,
        "operating_margin": op_margin,
        "net_margin": net_margin,
        "roe": roe,
        "revenue_growth": revenue_growth,
        "headcount": headcount,
        "sales_per_employee": sales_employee,
        "cfo": cfo,
        "fcf": fcf,
        "dividends_paid": dividends_paid,
        "buybacks": buybacks,
        "tcs_revenue_growth": tcs_growth,
        "tcs_operating_margin": tcs_margin,
        "tcs_roe": tcs_roe,
    },
    "geography_pct": geography,
    "segments_pct": segments,
    "infosys_monthly_price": infosys_price,
    "tcs_monthly_price": tcs_price,
    "comparative_returns": {
        "note": "Populate this section only from an authorized, redistributable comparison dataset.",
    },
    "metrics": metrics,
}

(OUT / "analysis_data.json").write_text(json.dumps(data, indent=2), encoding="utf-8")
print(json.dumps(metrics, indent=2))
