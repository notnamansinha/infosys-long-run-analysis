from pathlib import Path
from openpyxl import load_workbook
import sys

root = Path(__file__).resolve().parent
prefix = sys.argv[1]
row_nums = [int(x) for x in sys.argv[2:]]
path = next(root.glob(prefix + "*/*.xlsx"))
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
print(path.relative_to(root), ws.title, ws.max_row, ws.max_column)
for r in row_nums:
    vals = [ws.cell(r, c).value for c in range(1, (ws.max_column or 0) + 1)]
    print(r, vals)
